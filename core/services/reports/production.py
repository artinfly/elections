"""
Общее для отчётов «по производствам» (конструктор и формат ВН).

Описание:
    Раскладка людей по производствам и цехам, порядок производств, заголовок
    отчёта, ведущие колонки шапки и строка-название производства.
    Производство берётся из поля service — так исторически устроены все отчёты
    «по производствам» (см. reports/standard.py).
"""

from collections import defaultdict
from collections.abc import Iterable
from datetime import datetime
from typing import Any

from openpyxl.styles import Alignment, Font

from ...models import Employee
from ..excel import NO_PRODUCTION

# Шрифт и выравнивание заголовков и итоговых строк.
BOLD = Font(bold=True)
CENTERED = Alignment(horizontal="center", vertical="center", wrap_text=True)
# Выравнивание чисел в строках цехов и производств.
CENTER = Alignment(horizontal="center")

# Номер строки, с которой начинается двухстрочная шапка колонок.
HEADER_ROW = 2
# Первая строка данных (после заголовка отчёта и двух строк шапки).
FIRST_DATA_ROW = 4


def lead_headers(include_depts: bool) -> tuple[str, ...]:
    """Ведущие колонки: производство, (цех), всего людей, всего 19 округ."""
    if include_depts:
        return ("Производство", "Цех", "Всего", "Всего 19 округ")
    return ("Производство", "Всего", "Всего 19 округ")


def group_people(people: Iterable[Employee], include_depts: bool) -> dict[str, Any]:
    """
    Раскладывает людей по производствам (и цехам внутри них).

    Возвращает:
        {производство: [люди]} или {производство: {цех: [люди]}}.
        Люди без производства попадают в группу NO_PRODUCTION.
    """
    data = (
        defaultdict(lambda: defaultdict(list)) if include_depts else defaultdict(list)
    )
    for person in people:
        production = person.service or NO_PRODUCTION
        if include_depts:
            data[production][person.department].append(person)
        else:
            data[production].append(person)
    return data


def production_order(names: Iterable[str]) -> list[str]:
    """Производства по алфавиту, «Без производства» — всегда последним."""
    return sorted(names, key=lambda name: (name == NO_PRODUCTION, name))


def write_title(sheet: Any, moment: datetime, total_columns: int) -> None:
    """Заголовок отчёта в первой строке на всю ширину таблицы."""
    title = sheet.cell(
        1, 1, f"Итоговый отчёт по видам производства на {moment:%d.%m.%y %H:%M}"
    )
    title.font = Font(bold=True, size=12)
    title.alignment = CENTERED
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)


def draw_lead_headers(sheet: Any, titles: tuple[str, ...]) -> int:
    """
    Рисует ведущие колонки шапки (каждая объединена на две строки).

    Возвращает:
        int: номер следующей свободной колонки.
    """
    column = 1
    for title in titles:
        sheet.merge_cells(
            start_row=HEADER_ROW,
            start_column=column,
            end_row=HEADER_ROW + 1,
            end_column=column,
        )
        cell = sheet.cell(HEADER_ROW, column, title)
        cell.font = BOLD
        cell.alignment = CENTERED
        column += 1
    return column


def write_production_heading(
    sheet: Any, row: int, production: str, total_columns: int
) -> None:
    """Строка-название производства на всю ширину (в режиме с цехами)."""
    sheet.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=total_columns
    )
    cell = sheet.cell(row, 1, production)
    cell.font = BOLD
    cell.alignment = Alignment(horizontal="center")
