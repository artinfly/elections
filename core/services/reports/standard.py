"""
Модуль генерации стандартных Excel-отчётов и ZIP-архивов.

Описание:
    Содержит логику создания «бумажных» отчётов: разделение по производствам,
    способы голосования, сводные таблицы, полная выгрузка сотрудников,
    отчёты по конкретным цехам и упаковка этих отчётов в ZIP-архивы.
    Конструктор отчётов по фильтрам — в reports/custom.py.
"""

from datetime import datetime
from typing import Any, Callable, Optional

import openpyxl
from django.db.models import Count, Q
from django.utils import timezone
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ...models import DEG, METHOD_LABELS, UIK, UIK19, UVZ, Employee
from ..archiver import ReportArchiver
from ..excel import (
    COLUMNS,
    NO_PRODUCTION,
    REPORT_MODES,
    REPORT_WIDTHS,
    _by_number,
    department_file_name,
    padded_number,
)
from .custom import department_custom_report
from .production import BOLD, CENTERED, production_order

# ==============================================================================
# Вспомогательные функции
# ==============================================================================


def _unique_name(taken: set, name: str) -> str:
    """
    Генерирует уникальное имя файла, избегая дубликатов в архиве.

    Описание:
        Если имя уже занято, добавляет суффикс "-2", "-3" и так далее,
        пока не будет найдено свободное имя. Регистрирует результат в наборе.

    Аргументы:
        taken: набор уже занятых имён (изменяется внутри функции).
        name: желаемое имя файла.

    Возвращает:
        str: гарантированно уникальное имя.
    """
    unique, attempt = name, 2
    while unique in taken:
        unique = f"{name}-{attempt}"
        attempt += 1
    taken.add(unique)
    return unique


def _share_row(
    sheet: Any, line: int, label: str, values: dict, bold: bool = False
) -> int:
    """
    Строка отчёта «Разделение по производствам»: работающих, проголосовавших
    по QR-коду, через ответственного и процент явки.

    Аргументы:
        values: словарь с ключами people, came, proxy.
        bold: жирный шрифт (строки «Итого» и «Всего»).

    Возвращает:
        int: номер следующей свободной строки.
    """
    people, came = values["people"], values["came"]
    row = (label, people, came, values["proxy"], came / people if people else 0)
    _write_cells(sheet, line, row, bold)
    return line + 1


def _method_share_row(
    sheet: Any, line: int, label: str, values: dict, bold: bool = False
) -> int:
    """
    Строка отчёта «Способы по производствам»: число людей по каждому способу,
    сколько выбрали способ всего и их доля.

    Аргументы:
        values: словарь с ключами people, deg, uik, uvz, u19.
        bold: жирный шрифт (строки «Итого» и «Всего»).

    Возвращает:
        int: номер следующей свободной строки.
    """
    people = values["people"]
    by_method = [values[key] for key in ("deg", "uik", "uvz", "u19")]
    chosen = sum(by_method)
    row = (label, people, *by_method, chosen, chosen / people if people else 0)
    _write_cells(sheet, line, row, bold)
    return line + 1


def _write_cells(sheet: Any, line: int, values: tuple, bold: bool) -> None:
    """Пишет строку: числа по центру, последняя колонка — процент (0.00%)."""
    for column, value in enumerate(values, 1):
        cell = sheet.cell(line, column, value)
        if column > 1:
            cell.alignment = Alignment(horizontal="center")
        if bold:
            cell.font = BOLD
    sheet.cell(line, len(values)).number_format = "0.00%"


def _draw_header(sheet: Any, merges: tuple, titles: tuple) -> None:
    """Двухстрочная шапка: объединения ячеек и подписи (жирным, по центру)."""
    for cells in merges:
        sheet.merge_cells(cells)
    for coordinate, title in titles:
        cell = sheet[coordinate]
        cell.value = title
        cell.font = BOLD
        cell.alignment = CENTERED


def _departments_by_production(queryset: Any, **counts: Any) -> dict:
    """
    Цеха с подсчётами, сгруппированные по производствам (поле service).

    ВАЖНО: «Производство» в отчётах строится по полю service — историческое
    расхождение имён. Не переименовывать на production без синхронного
    изменения импорта из внешней базы.

    Возвращает:
        {производство: [строки цехов]} — производства в порядке production_order,
        цеха внутри — по номеру (числовые первыми).
    """
    grouped: dict = {}
    rows = queryset.values("service", "department").annotate(**counts).order_by()
    for row in rows:
        grouped.setdefault(row["service"] or NO_PRODUCTION, []).append(row)
    return {
        production: sorted(
            grouped[production], key=lambda r: _by_number(r["department"])
        )
        for production in production_order(grouped)
    }


def _write_production_blocks(
    sheet: Any, grouped: dict, keys: tuple, write_row: Any, width: int
) -> None:
    """
    Блоки по производствам с третьей строки: название производства на всю
    ширину таблицы, строки цехов, «Итого» по производству; после пустой
    строки — «Всего».
    """
    line = 3
    totals = dict.fromkeys(keys, 0)
    for production, rows in grouped.items():
        sheet.merge_cells(
            start_row=line, start_column=1, end_row=line, end_column=width
        )
        sheet.cell(line, 1, production).font = BOLD
        line += 1

        subtotal = dict.fromkeys(keys, 0)
        for row in rows:
            line = write_row(sheet, line, padded_number(row["department"]), row)
            for key in keys:
                subtotal[key] += row[key]
        line = write_row(sheet, line, "Итого", subtotal, bold=True)
        for key in keys:
            totals[key] += subtotal[key]

    write_row(sheet, line + 1, "Всего", totals, bold=True)


def _set_widths(sheet: Any, widths: tuple) -> None:
    """Ширины колонок слева направо."""
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


# ==============================================================================
# Отчёты по производствам
# ==============================================================================


def production_table() -> Any:
    """
    Отчёт «Разделение по производствам (Голосование)».

    Описание:
        Цеха сгруппированы по производствам. Для каждого цеха — число
        работающих, проголосовавших по QR-коду и через ответственного,
        процент явки. Под каждым производством «Итого», в конце — «Всего».

    Возвращает:
        Книга openpyxl с готовым отчётом.
    """
    grouped = _departments_by_production(
        Employee.objects.exclude(department=""),
        people=Count("id"),
        came=Count("id", filter=Q(voted=True)),
        proxy=Count("id", filter=Q(proxy_vote=True)),
    )

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "По производствам"
    _set_widths(sheet, (16, 14, 18, 18, 10))
    sheet.freeze_panes = "A3"
    _draw_header(
        sheet,
        ("A1:A2", "B1:B2", "C1:E1"),
        (
            ("A1", "Подразделение"),
            ("B1", "Общее число работающих"),
            ("C1", "Итог"),
            ("C2", "Количество проголосовавших (QR-код)"),
            ("D2", "Количество проголосовавших (Ответственный)"),
            ("E2", "%"),
        ),
    )
    keys = ("people", "came", "proxy")
    _write_production_blocks(sheet, grouped, keys, _share_row, width=5)
    return book


def production_method_table(exclude_u19: bool = False) -> Any:
    """
    Отчёт «Способы голосования по производствам».

    Описание:
        Как «Разделение по производствам», но вместо явки — распределение
        сотрудников по плановым способам голосования.

    Аргументы:
        exclude_u19: если True, сотрудники 19-го округа исключаются из отчёта.

    Возвращает:
        Книга openpyxl с готовым отчётом.
    """
    queryset = Employee.objects.exclude(department="")
    if exclude_u19:
        queryset = queryset.exclude(okrug="19")
    grouped = _departments_by_production(
        queryset,
        people=Count("id"),
        deg=Count("id", filter=Q(method=DEG)),
        uik=Count("id", filter=Q(method=UIK)),
        uvz=Count("id", filter=Q(method=UVZ)),
        u19=Count("id", filter=Q(method=UIK19)),
    )

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Способы (без 19)" if exclude_u19 else "Способы по производствам"
    _set_widths(sheet, (16, 10, 10, 12, 10, 16, 20, 12))
    sheet.freeze_panes = "A3"
    _draw_header(
        sheet,
        ("A1:A2", "B1:B2", "C1:F1", "G1:H1"),
        (
            ("A1", "Подразделение"),
            ("B1", "Общее количество"),
            ("C1", "Способ голосования"),
            ("C2", "ДЭГ"),
            ("D2", "УИК"),
            ("E2", "УИК-УВЗ"),
            ("F2", "УИК-19"),
            ("G1", "Итог"),
            ("G2", "Кол-во зарегистрированных"),
            ("H2", "%"),
        ),
    )
    keys = ("people", "deg", "uik", "uvz", "u19")
    _write_production_blocks(sheet, grouped, keys, _method_share_row, width=8)
    return book


# ==============================================================================
# Сводная таблица
# ==============================================================================


# Колонки сводной таблицы после названия группы (в этом порядке).
SUMMARY_KEYS = (
    "people",
    "plan_deg",
    "plan_uik",
    "plan_uvz",
    "plan_u19",
    "came",
    "proxy_vote",
)


def summary_table(
    group_field: str = "department",
    group_title: str = "Подразделение",
    exclude_u19: bool = False,
) -> Any:
    """
    Сводная таблица: строка на группу (по умолчанию — цех).

    Описание:
        Для каждой группы — число людей, план по способам голосования (method),
        проголосовавшие по факту (voted), через ответственного и процент явки.
        Колонки способов считаются по плану, «Проголосовавшие» — по факту,
        поэтому их суммы не обязаны совпадать. Внизу строка «Итого».

    Аргументы:
        group_field: поле модели для группировки (например, "department" или "uik").
        group_title: заголовок первой колонки.
        exclude_u19: если True, сотрудники 19-го округа исключаются из отчёта.

    Возвращает:
        Книга openpyxl со сводной таблицей.
    """
    queryset = Employee.objects.exclude(**{group_field: ""})
    if exclude_u19:
        queryset = queryset.exclude(okrug="19")
    # order_by() сбрасывает сортировку модели, иначе она попадёт в GROUP BY.
    rows = sorted(
        queryset.values(group_field)
        .annotate(
            people=Count("id"),
            plan_deg=Count("id", filter=Q(method=DEG)),
            plan_uik=Count("id", filter=Q(method=UIK)),
            plan_uvz=Count("id", filter=Q(method=UVZ)),
            plan_u19=Count("id", filter=Q(method=UIK19)),
            came=Count("id", filter=Q(voted=True)),
            proxy_vote=Count("id", filter=Q(proxy_vote=True)),
        )
        .order_by(),
        key=lambda row: _by_number(row[group_field]),
    )

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Сводка"
    _set_widths(sheet, (18, 16, 10, 10, 12, 16, 10, 14, 10))
    sheet.freeze_panes = "A2"

    headers = (
        group_title,
        "Количество людей",
        "ДЭГ",
        "УИК",
        "УИК-УВЗ",
        "УИК-19",
        "Проголосовавшие",
        "Через ответственного",
        "Процент",
    )
    for column, name in enumerate(headers, 1):
        cell = sheet.cell(1, column, name)
        cell.font = BOLD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    def share(values: dict) -> float:
        return values["came"] / values["people"] if values["people"] else 0

    totals = dict.fromkeys(SUMMARY_KEYS, 0)
    line = 2
    for row in rows:
        # Номера цехов дополняются до трёх цифр, прочие группы пишутся как есть.
        label = row[group_field]
        if group_field == "department":
            label = padded_number(label)
        _write_cells(
            sheet, line, (label, *(row[k] for k in SUMMARY_KEYS), share(row)), False
        )
        for key in SUMMARY_KEYS:
            totals[key] += row[key]
        line += 1

    _write_cells(
        sheet, line, ("Итого", *(totals[k] for k in SUMMARY_KEYS), share(totals)), True
    )
    return book


# ==============================================================================
# Полная выгрузка и отчёт по цеху
# ==============================================================================


def export_xlsx() -> Any:
    """
    Полная выгрузка всех сотрудников в один Excel-файл.

    Описание:
        Выгружает все поля из модели плюс информацию о плане, явке и месте голосования.
        Использует итератор для экономии оперативной памяти на больших базах.

    Возвращает:
        Книга openpyxl со всеми данными.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Сотрудники"

    # Заголовки колонок.
    sheet.append(
        list(COLUMNS)
        + [
            "Способ (план)",
            "Проголосовал (QR-код)",
            "Где голосовал",
            "Через ответственного",
        ]
    )
    fields = list(COLUMNS.values())

    # Итератор с размером пачки 2000 строк не держит всю таблицу в памяти.
    for person in (
        Employee.objects.all()
        .order_by("department", "surname", "name", "patronymic")
        .iterator(chunk_size=2000)
    ):
        row = [getattr(person, f) for f in fields]
        row.append(METHOD_LABELS.get(person.method, ""))
        row.append("да" if person.voted else "нет")
        row.append(METHOD_LABELS.get(person.voted_method, ""))
        row.append("да" if person.proxy_vote else "нет")
        sheet.append(row)
    return book


def department_report(
    department: str, moment: Optional[datetime] = None, mode: str = "turnout"
) -> Any:
    """
    Формирует отдельный Excel-отчёт по одному цеху.

    Описание:
        Содержит шапку с датой, общую статистику (всего/сделано/процент)
        и список тех, у кого нужная отметка отсутствует (не проголосовали
        или не выбрали способ).

    Аргументы:
        department: номер цеха.
        moment: момент времени для шапки отчёта (по умолчанию текущее).
        mode: режим отчёта из REPORT_MODES — "turnout" (явка) или "method" (способ).

    Возвращает:
        Книга openpyxl с отчётом по цеху.
    """
    rule = REPORT_MODES[mode]
    moment = moment or timezone.localtime()

    # Всего и выполненных — одним запросом.
    stats = Employee.objects.filter(department=department).aggregate(
        total=Count("id"),
        done=Count("id", filter=rule["done"]),
    )
    total = stats["total"]
    done = stats["done"]

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = f"Цех {department}"[:31]

    for index, width in enumerate(REPORT_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A6"

    bold = Font(bold=True)
    title = sheet.cell(1, 1, f"{rule['title']} {moment:%d.%m.%y %H:%M}")
    title.font = Font(bold=True, size=12)
    sheet.cell(2, 1, f"Цех {department}").font = bold

    for column, name in enumerate(
        ("Общее количество голосующих", rule["column"], "Процент"), 1
    ):
        cell = sheet.cell(4, column, name)
        cell.font = bold
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    sheet.cell(5, 1, total).alignment = Alignment(horizontal="center")
    sheet.cell(5, 2, done).alignment = Alignment(horizontal="center")
    share = sheet.cell(5, 3, (done / total) if total else 0)
    share.number_format = "0.00%"
    share.alignment = Alignment(horizontal="center")

    sheet.cell(7, 1, rule["list_title"]).font = bold

    # Список тех, у кого нужной отметки нет.
    row = 8
    for person in (
        Employee.objects.filter(department=department)
        .exclude(rule["done"])
        .order_by("surname", "name", "patronymic")
    ):
        sheet.cell(row, 1, person.tab_number)
        sheet.cell(row, 2, person.fio)
        row += 1

    return book


# ==============================================================================
# Архивы по цехам
# ==============================================================================


def _archive_by_department(build_book: Callable[[str], Any]) -> ReportArchiver:
    """
    ZIP-архив с отдельной книгой Excel на каждый цех.

    Описание:
        Цеха берутся из базы (без пустых) и сортируются по номеру. Имя файла —
        номер цеха из трёх цифр (007.xlsx); совпавшие имена разводятся
        суффиксом «-2», «-3».

    Аргументы:
        build_book: функция, которая строит книгу по номеру цеха.

    Возвращает:
        ReportArchiver: наполненный сборщик (байты архива — build_bytes()).
    """
    departments = (
        Employee.objects.exclude(department="")
        .values_list("department", flat=True)
        .distinct()
        .order_by()
    )
    archiver = ReportArchiver()
    taken: set = set()
    for department in sorted(departments, key=_by_number):
        unique = _unique_name(taken, department_file_name(department))
        archiver.add_workbook(build_book(department), f"{unique}.xlsx")
    return archiver


def reports_archive(
    moment: Optional[datetime] = None, mode: str = "turnout"
) -> ReportArchiver:
    """
    Архив отчётов по цехам: явка (mode="turnout") или выбор способа ("method").

    Аргументы:
        moment: момент времени для шапок отчётов (по умолчанию — сейчас).
        mode: режим отчёта из REPORT_MODES.
    """
    moment = moment or timezone.localtime()
    return _archive_by_department(lambda dep: department_report(dep, moment, mode))


def custom_reports_archive(params: dict) -> ReportArchiver:
    """
    Архив отчётов конструктора «по людям» по каждому цеху.

    Аргументы:
        params: фильтры конструктора (применяются к содержимому каждого файла).
    """
    return _archive_by_department(lambda dep: department_custom_report(dep, params))


def _responsible_book(department: str) -> Any:
    """
    Шаблон отметок «Голосование через ответственного» по одному цеху.

    Список сотрудников цеха; у тех, кому отметка уже поставлена, в последней
    колонке стоит «+». Заполненный файл загружается обратно на странице загрузки.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = f"Цех {department}"[:31]

    sheet.append(["Цех", "Таб.№", "ФИО", "Проголосовал (Если проголосовал то 1)"])
    for cell in sheet[1]:
        cell.font = BOLD
    _set_widths(sheet, (10, 12, 42, 12))
    sheet.freeze_panes = "A2"

    # iterator() не держит в памяти весь цех целиком.
    people = (
        Employee.objects.filter(department=department)
        .order_by("surname", "name", "patronymic")
        .iterator(chunk_size=2000)
    )
    for person in people:
        mark = "+" if person.proxy_vote else ""
        sheet.append([person.department, person.tab_number, person.fio, mark])
    return book


def responsible_marks_archive() -> ReportArchiver:
    """Архив шаблонов «Голосование через ответственного» по каждому цеху."""
    return _archive_by_department(_responsible_book)
