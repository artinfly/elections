"""
Модуль отчёта по сотрудникам УИК-УВЗ с оформленным заявлением.

Описание:
    Формирует список сотрудников, у которых стоит отметка
    "Заявление оформил" на участке УИК-УВЗ (method == UVZ и mark_uvz == True),
    но при этом нет отметки о голосовании (voted == False).

    Колонки отчёта: № | Цех | Таб№ | ФИО.
    Первая колонка — порядковый номер записи, начинается с единицы,
    сквозной по всему отчёту (итоговые строки не нумеруются).
    После строк каждого цеха добавляется подытог: объединённая строка
    "Итого по цеху NNN: X человек". В конце — строка "Всего" по всему отчёту.
    Сортировка: сначала цех (по числовому порядку, как в других отчётах),
    затем ФИО (фамилия, имя, отчество).
"""

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from .helpers import _apply_border, _by_number, padded_number
from .models import UVZ, Employee


def _people_word(count: int) -> str:
    """
    Возвращает правильную форму слова "человек" для числа.

    Описание:
        1 человек, 2-4 человека, 5 и более человек, с учётом исключений
        11-14 (одиннадцать человек и т.д.).

    Аргументы:
        count: количество людей.

    Возвращает:
        str: нужная форма слова.
    """
    if count % 10 == 1 and count % 100 != 11:
        return "человек"
    if count % 10 in (2, 3, 4) and count % 100 not in (12, 13, 14):
        return "человека"
    return "человек"


def uvz_statement_not_voted_table() -> Workbook:
    """
    Строит книгу со списком "заявление на УВЗ оформлено, но не проголосовал".

    Описание:
        Выборка: method == UVZ, mark_uvz == True, voted == False.
        Первая строка — заголовки колонок (жирные, по центру), далее по строке
        на сотрудника: порядковый номер (с единицы), цех (дополненный до трёх
        цифр), табельный номер, ФИО. После каждого цеха — объединённая строка
        подытога с количеством людей в цехе, в конце — строка "Всего".
        Шапка зафиксирована при прокрутке, в конце применяется рамка таблицы.

    Возвращает:
        Workbook: книга openpyxl с готовым отчётом.
    """
    people = Employee.objects.filter(method=UVZ, mark_uvz=True, voted=False)

    # Сортировка: цех по числовому порядку, затем ФИО.
    rows = sorted(
        people,
        key=lambda p: (_by_number(p.department), p.surname, p.name, p.patronymic),
    )

    # Группировка по цехам с сохранением порядка сортировки:
    # список пар (цех, [сотрудники цеха]).
    groups: list[tuple[str, list]] = []
    for person in rows:
        if not groups or groups[-1][0] != person.department:
            groups.append((person.department, []))
        groups[-1][1].append(person)

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "УВЗ не проголосовали"
    bold = Font(bold=True)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Заголовки колонок: первой идёт порядковый номер.
    for column, title in enumerate(("№", "Цех", "Таб№", "ФИО"), 1):
        cell = sheet.cell(1, column, title)
        cell.font = bold
        cell.alignment = centered

    # Ширины колонок.
    for index, width in enumerate((6, 10, 12, 42), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    # Фиксация шапки при прокрутке.
    sheet.freeze_panes = "A2"

    # Строки сотрудников и подытоги по цехам.
    row = 2
    number = 1
    total = 0
    for department, persons in groups:
        for person in persons:
            sheet.cell(row, 1, number).alignment = centered
            sheet.cell(row, 2, padded_number(department)).alignment = centered
            sheet.cell(row, 3, person.tab_number)
            sheet.cell(row, 4, person.fio)
            row += 1
            number += 1

        # Подытог по цеху: объединённая строка на всю ширину таблицы.
        count = len(persons)
        total += count
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = sheet.cell(
            row,
            1,
            f"Итого по цеху {padded_number(department)}: {count} {_people_word(count)}",
        )
        cell.font = bold
        cell.alignment = centered
        row += 1

    # Строка "Всего" по всему отчёту.
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = sheet.cell(row, 1, f"Всего не проголосовало: {total} {_people_word(total)}")
    cell.font = bold
    cell.alignment = centered

    _apply_border(sheet)
    return book
